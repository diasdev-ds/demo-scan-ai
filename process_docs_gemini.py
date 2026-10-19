import os
import json
import fitz  # PyMuPDF
from google import genai
from google.genai import types
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

API_KEY = os.environ.get("GEMINI_API_KEY", "")
client = genai.Client(api_key=API_KEY)
MODEL = "gemini-2.5-flash"

DOCS_DIR = os.environ.get("DOCS_DIR", "./docs")  # папка со сканами документов

PROMPT_NAKLADNAYA = """
Ты — система распознавания документов. Перед тобой скан накладной на русском или казахском языке.

Извлеки данные и верни ТОЛЬКО валидный JSON без пояснений:

{
  "тип": "накладная",
  "номер": "",
  "дата": "",
  "поставщик": { "название": "", "ИИН_БИН": "" },
  "получатель": { "название": "", "ИИН_БИН": "" },
  "товары": [
    { "название": "", "количество": "", "единица": "", "цена": "", "сумма": "" }
  ],
  "итого": "",
  "НДС": "",
  "итого_с_НДС": ""
}

Если поле не найдено — оставь пустую строку. Числа пиши как строки, сохраняя формат из документа.
"""

PROMPT_AVR = """
Ты — система распознавания документов. Перед тобой скан акта выполненных работ (АВР) на русском или казахском языке.

Извлеки данные и верни ТОЛЬКО валидный JSON без пояснений:

{
  "тип": "авр",
  "номер": "",
  "дата": "",
  "исполнитель": { "название": "", "ИИН_БИН": "" },
  "заказчик": { "название": "", "ИИН_БИН": "" },
  "работы": [
    { "описание": "", "количество": "", "единица": "", "цена": "", "сумма": "" }
  ],
  "итого": "",
  "НДС": "",
  "итого_с_НДС": ""
}

Если поле не найдено — оставь пустую строку. Числа пиши как строки, сохраняя формат из документа.
"""


def pdf_to_images(pdf_path: str) -> list[bytes]:
    doc = fitz.open(pdf_path)
    images = []
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        images.append(pix.tobytes("png"))
    doc.close()
    return images


def extract_data(pdf_path: str, doc_type: str) -> dict:
    print(f"  Обрабатываю: {os.path.basename(pdf_path)}")
    images = pdf_to_images(pdf_path)
    prompt = PROMPT_NAKLADNAYA if doc_type == "накладная" else PROMPT_AVR

    parts = [types.Part.from_text(text=prompt)]
    for img_bytes in images:
        parts.append(types.Part.from_bytes(data=img_bytes, mime_type="image/png"))

    response = client.models.generate_content(
        model=MODEL,
        contents=[types.Content(role="user", parts=parts)],
        config=types.GenerateContentConfig(
            temperature=0,
            response_mime_type="application/json",
        ),
    )

    raw = response.text.strip().replace("```json", "").replace("```", "").strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        print(f"  ОШИБКА парсинга JSON для {os.path.basename(pdf_path)}")
        return {"ошибка": raw, "файл": os.path.basename(pdf_path)}


def style_header(ws, row, cols, fill_color="1F4E79"):
    fill = PatternFill("solid", fgColor=fill_color)
    font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_border(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def save_to_excel(nakladnye: list, avr_list: list, output_path: str):
    wb = openpyxl.Workbook()

    # === Лист: Накладные ===
    ws1 = wb.active
    ws1.title = "Накладные"
    headers = ["Файл", "Номер", "Дата", "Поставщик", "БИН поставщика",
               "Получатель", "БИН получателя", "Товар", "Кол-во",
               "Единица", "Цена", "Сумма", "Итого", "НДС", "Итого с НДС"]
    widths = [20, 15, 12, 30, 18, 30, 18, 35, 10, 10, 15, 15, 15, 15, 15]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws1.append(headers)
    style_header(ws1, 1, len(headers))
    ws1.row_dimensions[1].height = 30

    row_num = 2
    for item in nakladnye:
        filename = item.get("файл", "")
        data = item.get("данные", {})
        if "ошибка" in data:
            ws1.append([filename, "ОШИБКА РАСПОЗНАВАНИЯ"] + [""] * (len(headers) - 2))
            row_num += 1
            continue
        товары = data.get("товары", [{}]) or [{}]
        start_row = row_num
        for т in товары:
            ws1.append([
                filename, data.get("номер", ""), data.get("дата", ""),
                data.get("поставщик", {}).get("название", ""),
                data.get("поставщик", {}).get("ИИН_БИН", ""),
                data.get("получатель", {}).get("название", ""),
                data.get("получатель", {}).get("ИИН_БИН", ""),
                т.get("название", ""), т.get("количество", ""), т.get("единица", ""),
                т.get("цена", ""), т.get("сумма", ""),
                data.get("итого", "") if товары.index(т) == 0 else "",
                data.get("НДС", "") if товары.index(т) == 0 else "",
                data.get("итого_с_НДС", "") if товары.index(т) == 0 else "",
            ])
            row_num += 1
        if row_num - start_row > 1:
            for col in [1, 2, 3, 4, 5, 6, 7, 13, 14, 15]:
                ws1.merge_cells(start_row=start_row, start_column=col, end_row=row_num - 1, end_column=col)
                ws1.cell(row=start_row, column=col).alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        zebra = PatternFill("solid", fgColor="D6E4F0") if (nakladnye.index(item) % 2 == 0) else PatternFill("solid", fgColor="FFFFFF")
        for r in range(start_row, row_num):
            for c in range(1, len(headers) + 1):
                ws1.cell(row=r, column=c).fill = zebra
                ws1.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=True)
    add_border(ws1, 1, row_num - 1, 1, len(headers))
    ws1.freeze_panes = "A2"

    # === Лист: АВР ===
    ws2 = wb.create_sheet("АВР")
    headers2 = ["Файл", "Номер", "Дата", "Исполнитель", "БИН исполнителя",
                "Заказчик", "БИН заказчика", "Описание работ", "Кол-во",
                "Единица", "Цена", "Сумма", "Итого", "НДС", "Итого с НДС"]
    widths2 = [20, 15, 12, 30, 18, 30, 18, 40, 10, 10, 15, 15, 15, 15, 15]
    for i, w in enumerate(widths2, 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2.append(headers2)
    style_header(ws2, 1, len(headers2), fill_color="2E4053")
    ws2.row_dimensions[1].height = 30

    row_num2 = 2
    for item in avr_list:
        filename = item.get("файл", "")
        data = item.get("данные", {})
        if "ошибка" in data:
            ws2.append([filename, "ОШИБКА РАСПОЗНАВАНИЯ"] + [""] * (len(headers2) - 2))
            row_num2 += 1
            continue
        работы = data.get("работы", [{}]) or [{}]
        start_row = row_num2
        for р in работы:
            ws2.append([
                filename, data.get("номер", ""), data.get("дата", ""),
                data.get("исполнитель", {}).get("название", ""),
                data.get("исполнитель", {}).get("ИИН_БИН", ""),
                data.get("заказчик", {}).get("название", ""),
                data.get("заказчик", {}).get("ИИН_БИН", ""),
                р.get("описание", ""), р.get("количество", ""), р.get("единица", ""),
                р.get("цена", ""), р.get("сумма", ""),
                data.get("итого", "") if работы.index(р) == 0 else "",
                data.get("НДС", "") if работы.index(р) == 0 else "",
                data.get("итого_с_НДС", "") if работы.index(р) == 0 else "",
            ])
            row_num2 += 1
        if row_num2 - start_row > 1:
            for col in [1, 2, 3, 4, 5, 6, 7, 13, 14, 15]:
                ws2.merge_cells(start_row=start_row, start_column=col, end_row=row_num2 - 1, end_column=col)
                ws2.cell(row=start_row, column=col).alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        zebra = PatternFill("solid", fgColor="E8DAEF") if (avr_list.index(item) % 2 == 0) else PatternFill("solid", fgColor="FFFFFF")
        for r in range(start_row, row_num2):
            for c in range(1, len(headers2) + 1):
                ws2.cell(row=r, column=c).fill = zebra
                ws2.cell(row=r, column=c).alignment = Alignment(vertical="center", wrap_text=True)
    add_border(ws2, 1, row_num2 - 1, 1, len(headers2))
    ws2.freeze_panes = "A2"

    wb.save(output_path)
    print(f"\n  Excel сохранён: {output_path}")


def main():
    if not API_KEY:
        print("ОШИБКА: Установи переменную окружения GEMINI_API_KEY")
        return

    print("=" * 50)
    print("  AI Обработка документов — SiteDev.kz")
    print("=" * 50)

    nakladnye_dir = os.path.join(DOCS_DIR, "накладные")
    avr_dir = os.path.join(DOCS_DIR, "авр")
    nakladnye_results, avr_results = [], []

    print("\n[1/2] Накладные:")
    for f in sorted(os.listdir(nakladnye_dir)):
        if f.lower().endswith(".pdf"):
            data = extract_data(os.path.join(nakladnye_dir, f), "накладная")
            nakladnye_results.append({"файл": f, "данные": data})

    print("\n[2/2] АВР:")
    for f in sorted(os.listdir(avr_dir)):
        if f.lower().endswith(".pdf"):
            data = extract_data(os.path.join(avr_dir, f), "авр")
            avr_results.append({"файл": f, "данные": data})

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(os.path.dirname(__file__), f"результат_{timestamp}.xlsx")

    print("\nФормирую Excel...")
    save_to_excel(nakladnye_results, avr_results, output_path)

    print("\nГотово! Обработано:")
    print(f"  Накладные: {len(nakladnye_results)} документов")
    print(f"  АВР:       {len(avr_results)} документов")
    print(f"  Файл:      {output_path}")


if __name__ == "__main__":
    main()
