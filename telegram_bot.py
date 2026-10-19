"""
Telegram-бот — точка входа системы.

Бухгалтер отправляет фото или PDF документа в чат →
бот распознаёт (Gemini) → отправляет в 1С (OData) → отвечает результатом + Excel.

Запуск:  TELEGRAM_TOKEN=... GEMINI_API_KEY=... python3 telegram_bot.py
"""
import os
import io
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from telegram import Update
from telegram.ext import (
    Application, MessageHandler, CommandHandler, filters, ContextTypes
)

import extractor
import sender_1c

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
log = logging.getLogger(__name__)

TOKEN = os.environ.get("TELEGRAM_TOKEN", "")


def make_excel(data: dict) -> io.BytesIO:
    """Распознанный документ -> красивый Excel в памяти."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Документ"

    тип = data.get("тип", "накладная")
    is_avr = тип == "авр"

    # Шапка документа
    header_fill = PatternFill("solid", fgColor="1F4E79")
    rows = [
        ("Тип документа", "Акт выполненных работ" if is_avr else "Накладная"),
        ("Номер", data.get("номер", "")),
        ("Дата", data.get("дата", "")),
    ]
    if is_avr:
        rows += [
            ("Исполнитель", data.get("исполнитель", {}).get("название", "")),
            ("БИН исполнителя", data.get("исполнитель", {}).get("ИИН_БИН", "")),
            ("Заказчик", data.get("заказчик", {}).get("название", "")),
        ]
    else:
        rows += [
            ("Поставщик", data.get("поставщик", {}).get("название", "")),
            ("БИН поставщика", data.get("поставщик", {}).get("ИИН_БИН", "")),
            ("Получатель", data.get("получатель", {}).get("название", "")),
        ]

    r = 1
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1

    r += 1
    cols = ["Описание" if is_avr else "Товар", "Кол-во", "Ед.", "Цена", "Сумма"]
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=r, column=c, value=name)
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    r += 1

    позиции = data.get("работы" if is_avr else "товары", [])
    for p in позиции:
        ws.cell(row=r, column=1, value=p.get("описание") or p.get("название", ""))
        ws.cell(row=r, column=2, value=p.get("количество", ""))
        ws.cell(row=r, column=3, value=p.get("единица", ""))
        ws.cell(row=r, column=4, value=p.get("цена", ""))
        ws.cell(row=r, column=5, value=p.get("сумма", ""))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Итого").font = Font(bold=True)
    ws.cell(row=r, column=5, value=data.get("итого", ""))
    ws.cell(row=r + 1, column=1, value="НДС").font = Font(bold=True)
    ws.cell(row=r + 1, column=5, value=data.get("НДС", ""))
    ws.cell(row=r + 2, column=1, value="Итого с НДС").font = Font(bold=True)
    ws.cell(row=r + 2, column=5, value=data.get("итого_с_НДС", ""))

    ws.column_dimensions["A"].width = 40
    for col in "BCDE":
        ws.column_dimensions[col].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 Бот распознавания документов SiteDev.kz\n\n"
        "Отправьте мне фото или PDF накладной / акта выполненных работ.\n"
        "Я распознаю данные, создам документ в 1С и пришлю Excel с результатом."
    )


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    status = await msg.reply_text("📥 Получил документ. Распознаю…")

    try:
        # 1. Скачиваем файл (фото или PDF)
        if msg.photo:
            tg_file = await msg.photo[-1].get_file()
            filename = "photo.jpg"
        elif msg.document:
            tg_file = await msg.document.get_file()
            filename = msg.document.file_name or "document.pdf"
        else:
            await status.edit_text("⚠️ Пришлите фото или PDF документа.")
            return

        data_bytes = bytes(await tg_file.download_as_bytearray())

        # 2. Распознаём
        await status.edit_text("🔍 Распознаю данные через AI…")
        result = extractor.extract(data_bytes, filename, doc_type="auto")

        if "ошибка" in result:
            await status.edit_text("❌ Не удалось распознать документ. Попробуйте более чёткое фото.")
            return

        тип = "Акт выполненных работ" if result.get("тип") == "авр" else "Накладная"

        # 3. Отправляем в 1С
        await status.edit_text("📤 Создаю документ в 1С…")
        try:
            resp_1c = sender_1c.send_to_1c(result)
            num = resp_1c.get("Number", "—")
            status_1c = f"✅ Документ создан в 1С (черновик), №{num}"
        except Exception as e:
            log.exception("Ошибка отправки в 1С")
            status_1c = f"⚠️ 1С недоступна ({str(e)[:60]}). Данные распознаны, Excel ниже."

        # 4. Формируем ответ
        итого = result.get("итого_с_НДС") or result.get("итого", "—")
        контрагент = (
            result.get("исполнитель", {}).get("название")
            or result.get("поставщик", {}).get("название", "—")
        )
        summary = (
            f"📄 {тип}\n"
            f"№ {result.get('номер', '—')} от {result.get('дата', '—')}\n"
            f"Контрагент: {контрагент}\n"
            f"Сумма: {итого} ₸\n\n"
            f"{status_1c}"
        )
        await status.edit_text(summary)

        # 5. Excel
        excel = make_excel(result)
        await msg.reply_document(
            document=excel,
            filename=f"{result.get('тип', 'doc')}_{result.get('номер', '')}.xlsx",
            caption="📊 Результат распознавания"
        )

    except Exception as e:
        log.exception("Ошибка обработки")
        await status.edit_text(f"❌ Ошибка: {str(e)[:100]}")


def main():
    if not TOKEN:
        print("ОШИБКА: задай TELEGRAM_TOKEN")
        return
    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.PHOTO | filters.Document.ALL, handle_document))
    print("Бот запущен. Открой Telegram и напиши боту /start")
    app.run_polling()


if __name__ == "__main__":
    main()
